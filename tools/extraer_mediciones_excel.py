"""Extrae los valores ingresados de la matriz 2026 sin evaluar sus fórmulas."""

from __future__ import annotations

import json
import sys
from datetime import timedelta
from pathlib import Path

import openpyxl


MESES_REGISTRADOS = range(1, 8)


def valor_numerico(valor: object, indicador: str) -> float:
    if indicador == "continuidad":
        if not isinstance(valor, timedelta):
            raise ValueError(f"Se esperaba un valor de hora para continuidad, se recibió {valor!r}.")
        return round(valor.total_seconds() / 3600, 4)
    if not isinstance(valor, (int, float)):
        raise ValueError(f"Se esperaba un valor numérico para presión, se recibió {valor!r}.")
    return round(float(valor), 4)


def extraer_hoja(hoja: openpyxl.worksheet.worksheet.Worksheet, indicador: str) -> list[dict]:
    registros = []
    for fila in range(6, 114):
        codigo = hoja.cell(fila, 7).value
        if not codigo:
            continue
        for mes in MESES_REGISTRADOS:
            valor = hoja.cell(fila, 7 + mes).value
            if valor is None:
                continue
            registros.append({
                "codigo_punto": str(codigo).strip(),
                "anio": 2026,
                "mes": mes,
                "indicador": indicador,
                "valor": valor_numerico(valor, indicador),
            })
    return registros


def extraer_conexiones(hoja: openpyxl.worksheet.worksheet.Worksheet) -> list[dict]:
    conexiones = []
    zona_actual = ""
    sector_actual = ""
    for fila in range(6, 114):
        if hoja.cell(fila, 4).value:
            sector_actual = str(hoja.cell(fila, 4).value).strip()
        if hoja.cell(fila, 5).value:
            zona_actual = str(hoja.cell(fila, 5).value).strip()
        if not sector_actual or not zona_actual:
            continue
        for mes in MESES_REGISTRADOS:
            cantidad = hoja.cell(fila, 45 + mes).value
            if cantidad is None:
                continue
            conexiones.append({
                "sector": sector_actual,
                "zona_presion": zona_actual,
                "anio": 2026,
                "mes": mes,
                "cantidad": int(cantidad),
            })
    return conexiones


def main() -> None:
    entrada = Path(sys.argv[1])
    salida = Path(sys.argv[2])
    libro = openpyxl.load_workbook(entrada, data_only=False, read_only=False)
    continuidad = extraer_hoja(libro["CONTINUIDAD"], "continuidad")
    presion = extraer_hoja(libro["PRESION"], "presion")
    conexiones = extraer_conexiones(libro["CONTINUIDAD"])

    if len(continuidad) != 252 or len(presion) != 252 or len(conexiones) != 252:
        raise ValueError("La matriz debe contener 252 valores por indicador y 252 registros de conexiones de enero a julio.")

    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(json.dumps({
        "origen": entrada.name,
        "registros": continuidad + presion,
        "conexiones": conexiones,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Mediciones extraídas: {len(continuidad)} de continuidad, {len(presion)} de presión y {len(conexiones)} de conexiones.")


if __name__ == "__main__":
    main()
